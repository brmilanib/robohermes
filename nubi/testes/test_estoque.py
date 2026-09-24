# -*- coding: utf-8 -*-
"""Comparação do estoque (UpSeller): entradas, saídas, zerados, voltaram, novos e removidos."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import estoque  # noqa: E402


def it(sku, atual, **k):
    return dict({"sku": sku, "titulo": sku.lower(), "atual": atual, "disponivel": atual}, **k)


def test_comparar():
    a = [it("A", 6), it("B", 12), it("C", 0), it("D", 3), it("R", 1)]
    b = [it("A", 2), it("B", 0), it("C", 4), it("D", 3), it("N", 10, subtotal=500)]
    d = estoque.comparar(a, b)
    assert d["n"] == {"entradas": 1, "saidas": 2, "zeraram": 1, "voltaram": 1, "novos": 1, "removidos": 1}
    assert d["unidades_entraram"] == 4 and d["unidades_sairam"] == 16
    assert d["saidas"][0]["sku"] == "B" and d["zeraram"][0]["sku"] == "B"
    assert d["totais"]["unidades"] == 19 and d["totais"]["valor"] == 500 and d["totais"]["zerados"] == 1
    assert "saldo −3 un." in " ".join(estoque.resumo_texto(d))


def test_primeira_foto():
    d = estoque.comparar([], [it("A", 1)])
    assert not d["tem_anterior"] and "Primeira foto" in estoque.resumo_texto(d)[1]




def test_planilha_gestor():
    import io
    import openpyxl
    xs = [it("1050009143", 3, titulo="Home Spray", custo_medio=70.996), it("ARMAF-MEGA-200", 12, titulo="Body", custo_medio=48.0),
          it("SEM-CUSTO", 0, titulo="X")]
    ws = openpyxl.load_workbook(io.BytesIO(estoque.gerar_gestor(xs))).active
    v = list(ws.iter_rows(values_only=True))
    assert ws.title == "Planilha1" and v[0] == tuple(estoque.GESTOR_COLUNAS)
    assert v[1] == ("1050009143", "1050009143", None, "Home Spray", 71, None, None)
    assert v[2][4] == 48 and v[3][4] is None and ws["A2"].data_type == "s"


if __name__ == "__main__":
    test_comparar()
    test_primeira_foto()
    test_planilha_gestor()
    print("ok")
