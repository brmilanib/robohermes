# -*- coding: utf-8 -*-
"""
Vendedores monitorados: o export de um vendedor seguido no Nubimetrics, mês fechado.
Uma linha por anúncio do vendedor: título, marca, vendas em $, unidades, preço médio,
tipo de publicação, FULL, catálogo, frete grátis, desconto, SKU, GTIN, estado...

O arquivo não diz de quem é nem de que mês: o nome do arquivo é o nome do vendedor
(AUMA_PERFUMARIA_P2.xlsx -> "AUMA PERFUMARIA P2", igual à coluna Vendedor do Explorador)
e o mês é escolhido na importação (ou vem no nome, se tiver AAAA-MM).

O cruzamento usa as outras duas funções:
  - Ranking mensal de marcas (mesmo mês): quanto o vendedor representa das vendas de
    cada marca no mercado, posição/tendência/saturação da marca e o status do B.I.;
  - Explorador de anúncios: para os GTINs das marcas monitoradas, preço médio e número de
    vendedores no mercado e a posição deste vendedor.
"""

import io
import re

import nubi

PADRAO_MES = re.compile(r"(20\d{2})[-_.](\d{2})")


class ErroVendedor(Exception):
    pass


def _num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _sim(v):
    return str(v or "").strip().lower() in ("sim", "s", "yes", "true", "1")


def nome_do_arquivo(nome):
    """AUMA_PERFUMARIA_P2.xlsx -> AUMA PERFUMARIA P2 (tira extensão, datas e '(1)')."""
    base = re.sub(r"\.(xlsx|xls|csv)$", "", nome or "", flags=re.I)
    base = PADRAO_MES.sub(" ", base)
    base = re.sub(r"\(\d+\)", " ", base)
    base = re.sub(r"[_]+", " ", base)
    return re.sub(r"\s+", " ", base).strip(" -").upper()


def mes_do_arquivo(nome):
    m = PADRAO_MES.search(nome or "")
    return f"{m.group(1)}-{m.group(2)}" if m else ""


def ler_vendedor(dados, nome_arquivo=""):
    """Lê o export do vendedor (.xlsx ou .csv). Devolve (linhas, vendedor sugerido, mês sugerido)."""
    brutas = []
    if (nome_arquivo or "").lower().endswith(".csv"):
        import pandas as pd
        texto = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                texto = dados.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sep = ";" if texto and texto.count(";") > texto.count(",") else ","
        df = pd.read_csv(io.StringIO(texto or ""), sep=sep, dtype=str, keep_default_na=False)
        brutas = [list(df.columns)] + df.values.tolist()
        aba = ""
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001
            raise ErroVendedor("não consegui abrir o arquivo como planilha (.xlsx)")
        ws = wb.worksheets[0]
        aba = ws.title or ""
        brutas = [list(r) for r in ws.iter_rows(values_only=True)]

    ini = next((i for i, r in enumerate(brutas[:20])
                if any(str(c or "").strip().lower() in ("título", "titulo") for c in r)
                and any("vendas" in str(c or "").lower() for c in r)), None)
    if ini is None:
        raise ErroVendedor("não parece o export de anúncios de um vendedor do Nubimetrics "
                           "(faltam as colunas Título e Vendas)")
    cab = [str(c or "").strip() for c in brutas[ini]]
    idx = {c.lower(): i for i, c in enumerate(cab)}

    def col(r, *nomes):
        for n in nomes:
            i = idx.get(n.lower())
            if i is not None and i < len(r):
                return r[i]
        return None

    # Contexto para GTINs colados (dois GTINs grudados): prefixos de empresa que aparecem
    # sozinhos no próprio arquivo, pesados por unidades.
    conhecidos = {}
    for r in brutas[ini + 1:]:
        g = re.sub(r"\D", "", str(col(r, "GTIN") or ""))
        if len(g) in (12, 13):
            conhecidos[g[:7]] = conhecidos.get(g[:7], 0) + int(_num(col(r, "Vendas em Unid.", "Unidades")) or 1)

    linhas = []
    for r in brutas[ini + 1:]:
        titulo = str(col(r, "Título", "Titulo") or "").strip()
        if not titulo:
            continue
        marca = str(col(r, "Marca") or "").strip()
        sku = str(col(r, "SKU") or "").strip()
        gtin = nubi.normalizar_gtin(col(r, "GTIN"), conhecidos)
        if not gtin:
            gtin = nubi.normalizar_gtin(sku, conhecidos)
        vendas = _num(col(r, "Vendas em $", "Vendas"))
        un = int(_num(col(r, "Vendas em Unid.", "Unidades vendidas", "Unidades")))
        preco = _num(col(r, "Preço Médio", "Preco Medio", "Preço"))
        linhas.append({
            "titulo": titulo, "marca": marca, "marca_chave": nubi.compacta(marca) if marca else "",
            "gtin": gtin, "sku": sku, "vendas": vendas, "unidades": un,
            "preco": preco or (vendas / un if un else 0.0),
            "tipo_pub": str(col(r, "Tipo de Publicação", "Tipo de Publicacao") or "").strip(),
            "full": _sim(col(r, "Fulfillment", "Full")), "catalogo": _sim(col(r, "Catálogo.", "Catálogo", "Catalogo")),
            "frete_gratis": _sim(col(r, "Com Frete grátis", "Com Frete gratis")),
            "desconto": _sim(col(r, "Com desconto")),
            "estado": str(col(r, "Estado") or "").strip(),
            "bruto": {c: ("" if v is None else str(v)) for c, v in zip(cab, r) if c},
        })
    if not linhas:
        raise ErroVendedor("o arquivo não tem nenhum anúncio")
    vendedor = nome_do_arquivo(nome_arquivo) or nome_do_arquivo(aba)
    return linhas, vendedor, mes_do_arquivo(nome_arquivo)


def _div(a, b):
    return float(a) / float(b) if a is not None and b else 0.0


def resumo(linhas):
    v = sum(l["vendas"] or 0 for l in linhas)
    u = sum(l["unidades"] or 0 for l in linhas)
    pv = lambda f: _div(sum(l["vendas"] or 0 for l in linhas if f(l)), v)
    return {
        "vendas": v, "unidades": u, "ticket": _div(v, u), "anuncios": len(linhas),
        "com_venda": sum(1 for l in linhas if (l["unidades"] or 0) > 0),
        "ativos": sum(1 for l in linhas if (l["estado"] or "").lower() == "active"),
        "marcas": len({l["marca_chave"] for l in linhas if l["marca_chave"]}),
        "produtos": len({l["gtin"] or l["titulo"] for l in linhas}),
        "pct_full": pv(lambda l: l["full"]), "pct_catalogo": pv(lambda l: l["catalogo"]),
        "pct_premium": pv(lambda l: (l["tipo_pub"] or "").lower().startswith("premium")),
        "pct_desconto": pv(lambda l: l["desconto"]),
    }


def _cresc(a, b):
    return (_div(b, a) - 1) if a else None


BOM_RANKING = {"Subindo forte", "Crescimento consistente", "Nova"}
RUIM_RANKING = {"Em queda", "Saiu do ranking"}


def analisar(linhas, anteriores=None, rk_mes=None, rk_bi=None, explorador=None, vendedor=""):
    """
    linhas/anteriores: anúncios do vendedor no mês e no mês anterior (ou None).
    rk_mes: linhas do ranking de marcas do mesmo mês ({marca_chave: linha}) ou None.
    rk_bi: resultado de ranking.bi() até o mês (status das marcas, entradas) ou None.
    explorador: {gtin: {"produto", "un", "fat", "vendedores", "preco_medio", "pos_vendedor", "periodo"}}.
    """
    rk_mes = rk_mes or {}
    bi_marcas = {m["marca_chave"]: m for m in (rk_bi or {}).get("marcas", [])}
    explorador = explorador or {}
    z = resumo(linhas)
    if anteriores is not None:
        za = resumo(anteriores)
        z.update({"ant": za, "var_vendas": _cresc(za["vendas"], z["vendas"]),
                  "var_unidades": _cresc(za["unidades"], z["unidades"]),
                  "var_ticket": _cresc(za["ticket"], z["ticket"])})

    # --- marcas
    def por_marca(ls):
        d = {}
        for l in ls:
            k = l["marca_chave"] or "(SEM MARCA)"
            x = d.setdefault(k, {"marca": l["marca"] or "(sem marca)", "vendas": 0.0, "unidades": 0, "anuncios": 0,
                                 "full_v": 0.0, "cat_v": 0.0})
            x["vendas"] += l["vendas"] or 0
            x["unidades"] += l["unidades"] or 0
            x["anuncios"] += 1
            x["full_v"] += (l["vendas"] or 0) if l["full"] else 0
            x["cat_v"] += (l["vendas"] or 0) if l["catalogo"] else 0
        return d
    atual = por_marca(linhas)
    ant = por_marca(anteriores) if anteriores is not None else {}
    marcas, acum = [], 0.0
    for k, x in sorted(atual.items(), key=lambda kv: -kv[1]["vendas"]):
        rk = rk_mes.get(k)
        bi = bi_marcas.get(k)
        a = ant.get(k)
        share = _div(x["vendas"], z["vendas"])
        acum += share
        marcas.append({
            "marca": x["marca"], "marca_chave": k, "vendas": x["vendas"], "unidades": x["unidades"],
            "anuncios": x["anuncios"], "preco_medio": _div(x["vendas"], x["unidades"]),
            "share_vendedor": share, "share_acum": acum,
            "pct_full": _div(x["full_v"], x["vendas"]), "pct_catalogo": _div(x["cat_v"], x["vendas"]),
            "vendas_ant": a["vendas"] if a else None,
            "var_vendas": _cresc(a["vendas"], x["vendas"]) if a else None,
            "nova": anteriores is not None and a is None,
            "rk_posicao": rk["posicao"] if rk else None, "rk_vendas": rk["vendas"] if rk else None,
            "share_no_mercado": _div(x["vendas"], rk["vendas"]) if rk and rk["vendas"] else None,
            "rk_tendencia": rk["tendencia"] if rk else None, "rk_saturacao": rk["saturacao"] if rk else None,
            "rk_vendedores": rk["vendedores"] if rk else None,
            "rk_status": bi["status"] if bi else ("Fora do ranking" if rk_mes else None),
            "rk_nota": bi["nota"] if bi else None,
            "rk_cresc_3m": bi["cresc_3m"] if bi else None,
        })
    saiu = [{"marca": a["marca"], "marca_chave": k, "vendas_ant": a["vendas"], "unidades_ant": a["unidades"]}
            for k, a in sorted(ant.items(), key=lambda kv: -kv[1]["vendas"]) if k not in atual]

    # --- produtos (por GTIN; sem GTIN, pelo título)
    prods = {}
    for l in linhas:
        k = l["gtin"] or ("T:" + l["titulo"])
        p = prods.setdefault(k, {"gtin": l["gtin"], "marca": l["marca"], "vendas": 0.0, "unidades": 0, "anuncios": 0,
                                 "titulos": {}, "full": False, "catalogo": False})
        p["vendas"] += l["vendas"] or 0
        p["unidades"] += l["unidades"] or 0
        p["anuncios"] += 1
        p["full"] |= l["full"]
        p["catalogo"] |= l["catalogo"]
        p["titulos"][l["titulo"]] = p["titulos"].get(l["titulo"], 0) + (l["unidades"] or 0) + 1
    ant_prod = {}
    for l in anteriores or []:
        k = l["gtin"] or ("T:" + l["titulo"])
        ant_prod[k] = ant_prod.get(k, 0) + (l["vendas"] or 0)
    produtos = []
    for k, p in sorted(prods.items(), key=lambda kv: -kv[1]["vendas"]):
        ex = explorador.get(p["gtin"]) if p["gtin"] else None
        preco = _div(p["vendas"], p["unidades"])
        produtos.append({
            "gtin": p["gtin"], "produto": (ex or {}).get("produto") or max(p["titulos"], key=p["titulos"].get),
            "marca": p["marca"], "vendas": p["vendas"], "unidades": p["unidades"], "preco": preco,
            "anuncios": p["anuncios"], "full": "Sim" if p["full"] else "Não", "catalogo": "Sim" if p["catalogo"] else "Não",
            "share_vendedor": _div(p["vendas"], z["vendas"]),
            "var_vendas": _cresc(ant_prod.get(k), p["vendas"]) if anteriores is not None and ant_prod.get(k) else None,
            "novo": "Sim" if anteriores is not None and k not in ant_prod else "",
            "ex_preco_medio": ex["preco_medio"] if ex else None,
            "ex_dif_preco": (_div(preco, ex["preco_medio"]) - 1) if ex and ex["preco_medio"] and preco else None,
            "ex_vendedores": ex["vendedores"] if ex else None,
            "ex_pos_vendedor": ex["pos_vendedor"] if ex else None,
            "ex_share": ex["share_vendedor"] if ex else None,
            "ex_periodo": ex["periodo"] if ex else "",
        })

    # --- cruzamento com o ranking: oportunidades e riscos
    oportunidades, riscos = [], []
    vend_por_marca = {m["marca_chave"]: m for m in marcas}
    for b in (rk_bi or {}).get("marcas", []):
        if not b.get("no_ultimo"):
            continue
        m = vend_por_marca.get(b["marca_chave"])
        share_merc = m["share_no_mercado"] if m else 0
        if b["status"] in BOM_RANKING and (share_merc or 0) < 0.01:
            oportunidades.append({
                "marca": b["marca"], "status": b["status"], "nota": b["nota"], "posicao": b["posicao"],
                "vendas_mercado": b["vendas_atual"], "cresc_3m": b["cresc_3m"], "saturacao": b["saturacao"],
                "vendas_vendedor": m["vendas"] if m else 0, "share_no_mercado": share_merc or 0})
    oportunidades.sort(key=lambda x: -(x["nota"] or 0))
    for m in marcas:
        if m["share_vendedor"] >= 0.02 and (m["rk_status"] in RUIM_RANKING or
                                            (m["rk_cresc_3m"] is not None and m["rk_cresc_3m"] <= -0.15)):
            riscos.append(m)
    # Entradas do ranking ("vindo de baixo") que o vendedor ainda não trabalha
    vindo = {}
    for e in (rk_bi or {}).get("entradas", []):    # ordenadas por mês: fica a entrada mais recente
        if e.get("leitura") in ("Ficou e cresceu", "Entrou agora") and e["marca_chave"] not in vend_por_marca:
            vindo[e["marca_chave"]] = {"marca": e["marca"], "mes": e["mes"], "posicao": e["posicao"],
                                       "vendas": e["vendas"], "posicao_atual": e.get("posicao_atual"),
                                       "vendas_atual": e.get("vendas_atual"), "leitura": e["leitura"]}
    vindo = sorted(vindo.values(), key=lambda x: -(x["vendas_atual"] or x["vendas"] or 0))

    anuncios = [{k: v for k, v in l.items() if k not in ("bruto", "marca_chave")} for l in linhas]
    for a in anuncios:
        a["full"] = "Sim" if a["full"] else "Não"
        a["catalogo"] = "Sim" if a["catalogo"] else "Não"
        a["frete_gratis"] = "Sim" if a["frete_gratis"] else "Não"
        a["desconto"] = "Sim" if a["desconto"] else "Não"
    return {"resumo": z, "marcas": marcas, "marcas_sairam": saiu, "produtos": produtos,
            "oportunidades": oportunidades[:30], "riscos": riscos, "vindo_de_baixo": vindo[:30],
            "anuncios": anuncios}


def cruzar_explorador(anuncios_por_marca, vendedor):
    """
    anuncios_por_marca: {marca: (df de anúncios do período mais recente, "dd/mm a dd/mm")}.
    Devolve {gtin: {...}} com o mercado do GTIN no Explorador e a posição do vendedor.
    """
    alvo = nubi.compacta(vendedor)
    out = {}
    for marca, (df, periodo) in anuncios_por_marca.items():
        if df is None or df.empty:
            continue
        d = df[df["gtin"] != ""]
        for gtin, g in d.groupby("gtin"):
            un = int(g["un"].sum())
            fat = float(g["fat"].sum())
            por_v = g.groupby("vendedor")["un"].sum().sort_values(ascending=False)
            pos = None
            for i, v in enumerate(por_v.index, 1):
                if nubi.compacta(str(v)) == alvo:
                    pos = i
                    break
            un_v = int(g.loc[g["vendedor"].map(lambda v: nubi.compacta(str(v)) == alvo), "un"].sum())
            out[gtin] = {"produto": g["produto"].mode().iloc[0] if "produto" in g and not g["produto"].empty else "",
                         "un": un, "fat": fat, "vendedores": int((por_v > 0).sum()) or len(por_v),
                         "preco_medio": _div(fat, un), "pos_vendedor": pos, "share_vendedor": _div(un_v, un),
                         "periodo": periodo}
    return out


# ---------------------------------------------------------------------------
# Identidade do vendedor
# ---------------------------------------------------------------------------
# O Nubimetrics troca o nome de vendedores sem apelido por nomes aleatórios
# (BANTENG.PRETO.DEMONSTRATIVO), então o nome não é uma chave confiável.
#   Chave 1: o hash do vendedor no Nubimetrics (seller=<128 hex>), enviado pelo coletor.
#   Chave 2 (reserva): a "impressão digital" dos anúncios — os itens que mais vendem.

OFUSCADO = re.compile(r"^[A-Z]+[.\-][A-Z]+[.\-][A-Z]+$")          # BANTENG.PRETO.DEMONSTRATIVO / BANTENG-PRETO-...
LIMIAR_AUTO, LIMIAR_REVISAR = 0.60, 0.35


def ofuscado(nome):
    return bool(OFUSCADO.match((nome or "").strip().upper()))


def impressao(linhas, n=200):
    """Os ~200 itens que mais vendem (GTIN, senão SKU, senão título + marca), o total e o mix de marcas."""
    top = sorted(linhas, key=lambda l: -(l["vendas"] or 0))[:n]
    itens = set()
    for l in top:
        if l.get("gtin"):
            itens.add(l["gtin"])
        elif l.get("sku"):
            itens.add("S:" + str(l["sku"]).strip().upper())
        else:
            itens.add("T:" + nubi.normalizar(l["titulo"]) + "|" + nubi.compacta(l.get("marca") or ""))
    total = sum(l["vendas"] or 0 for l in linhas)
    marcas = {}
    for l in linhas:
        k = l.get("marca_chave") or ""
        marcas[k] = marcas.get(k, 0) + (l["vendas"] or 0)
    top_m = dict(sorted(marcas.items(), key=lambda kv: -kv[1])[:30])
    return {"itens": sorted(itens), "vendas": total,
            "marcas": {k: round(v / total, 4) for k, v in top_m.items()} if total else {}}


def comparar(a, b):
    """Similaridade entre duas impressões: itens em comum ÷ itens do menor conjunto,
    razão de faturamento (menor ÷ maior) e parecença do mix de marcas (cosseno)."""
    ia, ib = set(a.get("itens") or []), set(b.get("itens") or [])
    sim = len(ia & ib) / min(len(ia), len(ib)) if ia and ib else 0.0
    va, vb = a.get("vendas") or 0, b.get("vendas") or 0
    razao = min(va, vb) / max(va, vb) if va and vb else 0.0
    ma, mb = a.get("marcas") or {}, b.get("marcas") or {}
    num = sum(ma[k] * mb.get(k, 0) for k in ma)
    den = (sum(v * v for v in ma.values()) ** 0.5) * (sum(v * v for v in mb.values()) ** 0.5)
    tam = min(len(ia), len(ib)) / max(len(ia), len(ib)) if ia and ib else 0.0
    return {"itens": sim, "vendas": razao, "marcas": num / den if den else 0.0, "tamanho": tam}


def decidir(c):
    """Mesmo vendedor só com itens E faturamento E mix de marcas E tamanho de catálogo parecidos:
    lojas de perfume vendem muitos GTINs em comum, então itens sozinhos não bastam."""
    # catálogo pequeno "cabe" inteiro no top 200 de um grande: tamanhos parecidos também são exigidos
    if c["itens"] >= LIMIAR_AUTO and c["vendas"] >= 0.5 and c["marcas"] >= 0.8 and c.get("tamanho", 1) >= 0.6:
        return "mesmo"
    if c["itens"] >= LIMIAR_REVISAR:
        return "revisar"
    return "novo"
