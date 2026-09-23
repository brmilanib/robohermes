# -*- coding: utf-8 -*-
"""
Ranking mensal de marcas: o relatório "MARCAS" do Nubimetrics (fechamento do mês,
sem produtos). Uma linha por marca: posição, variação no ranking, vendas em $,
quantidade vendida, tendência, % catálogo, vendedores, saturação e ranking de demanda.

Nome do arquivo no padrão do Nubimetrics:  MARCAS-MLB1246-MLB6284-2026-08-01_1.xlsx
  -> categoria "MLB1246-MLB6284" (Beleza e Cuidado Pessoal > Perfumes), mês 2026-08.
"""

import hashlib
import io
import math
import re
import statistics

import nubi

PADRAO_ARQUIVO = re.compile(r"MARCAS[-_]+(?P<cat>(?:MLB\d+[-_]?)+?)[-_]+(?P<mes>\d{4}-\d{2})(?:-\d{2})?", re.I)

# Nomes das categorias do Mercado Livre que aparecem nos códigos (os que não estão aqui
# aparecem pelo código mesmo).
NOMES_CATEGORIA = {
    "MLB1246": "Beleza e Cuidado Pessoal",
    "MLB6284": "Perfumes",
    "MLB1248": "Maquiagem",
    "MLB1263": "Cuidados com a Pele",
    "MLB199407": "Cuidados com o Cabelo",
}

COLUNAS = ["#", "Variação", "Marca", "Vendas em $", "Quantidade vendas", "Tendencia ranking",
           "Catálogo", "Vendedores", "Saturação", "Ranking de demanda"]

MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
         "setembro", "outubro", "novembro", "dezembro"]


class ErroRanking(Exception):
    pass


def nome_categoria(cat):
    partes = [p for p in re.split(r"[-_]", cat or "") if p]
    return " > ".join(NOMES_CATEGORIA.get(p.upper(), p.upper()) for p in partes)


def nome_mes(mes):
    """'2026-08-01' -> 'agosto de 2026'."""
    a, m = str(mes)[:7].split("-")
    return f"{MESES[int(m) - 1]} de {a}"


def valor_abreviado(v):
    """'+$20.9M' -> 20900000.0 ; '+161.6k' -> 161600.0 ; '61,3%' -> 0.613 ; 549 -> 549."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("R$", "").replace("$", "").replace("+", "").replace(" ", "")
    pct = t.endswith("%")
    t = t.rstrip("%")
    m = re.fullmatch(r"(-?[\d.,]+)([kKmMbB]?)", t)
    if not m:
        return None
    num, suf = m.groups()
    if suf:                                   # abreviado: ponto é decimal (20.9M)
        num = num.replace(",", ".")
    elif "," in num:                          # pt-BR: 1.234,5
        num = num.replace(".", "").replace(",", ".")
    try:
        x = float(num)
    except ValueError:
        return None
    x *= {"k": 1e3, "m": 1e6, "b": 1e9}.get(suf.lower(), 1) if suf else 1
    return x / 100 if pct else x


def _inteiro(v):
    x = valor_abreviado(v)
    return None if x is None else int(round(x))


def ler_relatorio(dados, nome_arquivo=""):
    """Lê o .xlsx (ou .csv) do relatório de marcas. Devolve (linhas, categoria, mes)."""
    linhas_brutas = []
    if nome_arquivo.lower().endswith(".csv"):
        import pandas as pd
        texto = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                texto = dados.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sep = ";" if texto and texto.count(";") > texto.count(",") else ","
        df = pd.read_csv(io.StringIO(texto or ""), sep=sep, dtype=str, keep_default_na=False)
        linhas_brutas = [list(df.columns)] + df.values.tolist()
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001
            raise ErroRanking("não consegui abrir o arquivo como planilha (.xlsx)")
        ws = wb.worksheets[0]
        linhas_brutas = [list(r) for r in ws.iter_rows(values_only=True)]

    # Cabeçalho: a primeira linha que tem "Marca" e "Vendas".
    ini = next((i for i, r in enumerate(linhas_brutas[:20])
                if any(str(c or "").strip().lower() == "marca" for c in r)
                and any("venda" in str(c or "").lower() for c in r)), None)
    if ini is None:
        raise ErroRanking("não parece o relatório de marcas do Nubimetrics (faltam as colunas Marca e Vendas)")
    cab = [str(c or "").strip() for c in linhas_brutas[ini]]
    idx = {c.lower(): i for i, c in enumerate(cab)}

    def col(r, *nomes):
        for n in nomes:
            i = idx.get(n.lower())
            if i is not None and i < len(r):
                return r[i]
        return None

    linhas = []
    for r in linhas_brutas[ini + 1:]:
        marca = str(col(r, "Marca") or "").strip()
        if not marca:
            continue
        linhas.append({
            "posicao": _inteiro(col(r, "#", "Posição", "Posicao")),
            "variacao": _inteiro(col(r, "Variação", "Variacao")),
            "marca": marca,
            "marca_chave": nubi.compacta(marca),
            "vendas": valor_abreviado(col(r, "Vendas em $", "Vendas")),
            "unidades": valor_abreviado(col(r, "Quantidade vendas", "Unidades vendidas", "Quantidade")),
            "tendencia": str(col(r, "Tendencia ranking", "Tendência ranking", "Tendencia") or "").strip(),
            "catalogo": valor_abreviado(col(r, "Catálogo", "Catalogo")),
            "vendedores": _inteiro(col(r, "Vendedores")),
            "saturacao": str(col(r, "Saturação", "Saturacao") or "").strip(),
            "ranking_demanda": _inteiro(col(r, "Ranking de demanda")),
            "bruto": {c: ("" if v is None else str(v)) for c, v in zip(cab, r) if c},
        })
    if not linhas:
        raise ErroRanking("o arquivo não tem nenhuma marca")
    for i, l in enumerate(linhas, 1):
        if l["posicao"] is None:
            l["posicao"] = i

    m = PADRAO_ARQUIVO.search(nome_arquivo or "")
    categoria = m.group("cat").strip("-_").upper().replace("_", "-") if m else ""
    mes = f"{m.group('mes')}-01" if m else ""
    return linhas, categoria, mes


def hash_de(dados):
    return hashlib.sha256(dados).hexdigest()


# ---------------------------------------------------------------------------
# Análise do mês (valores calculados para a página)
# ---------------------------------------------------------------------------

def _div(a, b):
    return float(a) / float(b) if a is not None and b else 0.0


SAT_LIVRE = {"baixa": 1.0, "média": 0.5, "media": 0.5, "alta": 0.0}
TEND = {"crescendo": 1.0, "estável": 0.5, "estavel": 0.5, "diminuindo": 0.0}


def analisar(linhas, anteriores=None, marcas_explorador=()):
    """
    Enriquece as linhas do mês:
      ticket médio (vendas ÷ unidades), vendas por vendedor, share do top listado,
      comparação com o mês anterior (se houver) e nota/sinais de oportunidade.
    Nota de oportunidade (0–100), pensada para achar marca para entrar, não a maior:
      100 × √(tamanho, com teto nas 25% maiores) × (0,10 + 0,30 × saturação livre
      + 0,30 × tendência + 0,15 × vendas por vendedor [3× a mediana = 1]
      + 0,15 × subida no ranking [+5 posições = 1]).
    Acima do percentil 75 de vendas o tamanho não soma mais: a partir daí pesam
    crescimento, saturação e concorrência.
    """
    tot_v = sum(l["vendas"] or 0 for l in linhas)
    tot_u = sum(l["unidades"] or 0 for l in linhas)
    vendas_ord = sorted(l["vendas"] or 0 for l in linhas)
    p75 = vendas_ord[int(len(vendas_ord) * 0.75)] if vendas_ord else 1
    vpv = [_div(l["vendas"], l["vendedores"]) for l in linhas if l["vendedores"]]
    med_vpv = statistics.median(vpv) if vpv else 0
    ant = {l["marca_chave"]: l for l in (anteriores or [])}
    explorador = {nubi.compacta(m): m for m in marcas_explorador}
    saida, acum = [], 0.0
    for l in sorted(linhas, key=lambda x: x["posicao"] or 9999):
        x = {k: v for k, v in l.items() if k != "bruto"}
        x["ticket"] = _div(l["vendas"], l["unidades"])
        x["vendas_por_vendedor"] = _div(l["vendas"], l["vendedores"])
        x["share"] = _div(l["vendas"], tot_v)
        acum += x["share"]
        x["share_acum"] = acum
        a = ant.get(l["marca_chave"])
        if anteriores is not None:
            x["posicao_ant"] = a["posicao"] if a else None
            x["vendas_ant"] = a["vendas"] if a else None
            x["unidades_ant"] = a["unidades"] if a else None
            x["var_vendas"] = (_div(l["vendas"], a["vendas"]) - 1) if a and a["vendas"] else None
            x["var_unidades"] = (_div(l["unidades"], a["unidades"]) - 1) if a and a["unidades"] else None
            x["novo"] = a is None
        sat = SAT_LIVRE.get((l["saturacao"] or "").lower(), 0.5)
        ten = TEND.get((l["tendencia"] or "").lower(), 0.5)
        rel = min(1.0, _div(x["vendas_por_vendedor"], 3 * med_vpv)) if med_vpv else 0
        sub = min(1.0, max(0, l["variacao"] or 0) / 5)
        tam = math.sqrt(min(1.0, _div(l["vendas"] or 0, p75)))
        x["nota"] = round(100 * tam * (0.1 + 0.3 * sat + 0.3 * ten + 0.15 * rel + 0.15 * sub))
        s = []
        if (l["saturacao"] or "").lower() == "baixa":
            s.append("Pouca saturação")
        if (l["tendencia"] or "").lower() == "crescendo":
            s.append("Crescendo")
        if med_vpv and x["vendas_por_vendedor"] >= 3 * med_vpv:
            s.append("Poucos vendedores p/ o volume")
        if (l["variacao"] or 0) >= 3:
            s.append("Subiu no ranking")
        if l["catalogo"] is not None and l["catalogo"] < 0.5:
            s.append("Catálogo pouco disputado")
        if (l["tendencia"] or "").lower() == "diminuindo":
            s.append("Diminuindo")
        if (l["variacao"] or 0) <= -3:
            s.append("Caiu no ranking")
        if (l["saturacao"] or "").lower() == "alta":
            s.append("Saturada")
        if anteriores is not None and x.get("novo"):
            s.append("Nova no ranking")
        x["sinais"] = " · ".join(s)
        x["explorador"] = explorador.get(l["marca_chave"])
        saida.append(x)
    saiu = []
    if anteriores is not None:
        atuais = {l["marca_chave"] for l in linhas}
        saiu = [{"marca": a["marca"], "posicao_ant": a["posicao"], "vendas_ant": a["vendas"]}
                for a in anteriores if a["marca_chave"] not in atuais]
    conc = {}
    for n in (1, 3, 5, 10, 20):
        conc[f"top{n}"] = sum(x["share"] for x in saida[:n])
    resumo = {
        "marcas": len(saida), "vendas": tot_v, "unidades": tot_u, "ticket": _div(tot_v, tot_u),
        "vendedores_mediana": statistics.median([l["vendedores"] for l in linhas if l["vendedores"]] or [0]),
        "crescendo": sum(1 for l in linhas if (l["tendencia"] or "").lower() == "crescendo"),
        "diminuindo": sum(1 for l in linhas if (l["tendencia"] or "").lower() == "diminuindo"),
        "estavel": sum(1 for l in linhas if (l["tendencia"] or "").lower() in ("estável", "estavel")),
        "sat_baixa": sum(1 for l in linhas if (l["saturacao"] or "").lower() == "baixa"),
        "sat_media": sum(1 for l in linhas if (l["saturacao"] or "").lower() in ("média", "media")),
        "sat_alta": sum(1 for l in linhas if (l["saturacao"] or "").lower() == "alta"),
        "concentracao": conc,
    }
    if anteriores is not None:
        tv_a = sum(a["vendas"] or 0 for a in anteriores)
        tu_a = sum(a["unidades"] or 0 for a in anteriores)
        resumo.update({"vendas_ant": tv_a, "unidades_ant": tu_a,
                       "var_vendas": (_div(tot_v, tv_a) - 1) if tv_a else None,
                       "var_unidades": (_div(tot_u, tu_a) - 1) if tu_a else None,
                       "novas": sum(1 for x in saida if x.get("novo")), "sairam": len(saiu)})
    return saida, saiu, resumo


# ---------------------------------------------------------------------------
# B.I.: todos os meses importados de uma categoria, lado a lado
# ---------------------------------------------------------------------------

def _clip(v, lo=0.0, hi=1.0):
    return max(lo, min(hi, v))


def _cresc(a, b):
    """Variação de a para b (None se não dá para comparar)."""
    return (_div(b, a) - 1) if a and b is not None else None


def bi(meses, linhas_por_mes, marcas_explorador=()):
    """
    meses: lista ordenada de 'AAAA-MM-01'; linhas_por_mes: lista paralela com as linhas de cada mês.
    O relatório traz só as maiores marcas do mês: marca ausente num mês = fora do ranking
    (não quer dizer venda zero), por isso crescimento só compara meses em que a marca aparece.

    Nota B.I. (0–100), para achar marca com momento bom e sustentado:
      100 × √(tamanho no último mês, teto nas 25% maiores) × (0,30 × crescimento em 3 meses [+50% = 1]
      + 0,20 × consistência [meses em alta ÷ comparações] + 0,15 × subida no ranking em 3 meses [+10 = 1]
      + 0,20 × saturação livre + 0,15 × vendas crescendo mais que vendedores [+50 p.p. = 1]).
    """
    n = len(meses)
    resumo_meses = []
    ant = None
    for mes, linhas in zip(meses, linhas_por_mes):
        _, saiu, z = analisar(linhas, ant)
        vpv = [_div(l["vendas"], l["vendedores"]) for l in linhas if l["vendedores"]]
        resumo_meses.append({
            "mes": mes[:7], "nome": nome_mes(mes), "vendas": z["vendas"], "unidades": z["unidades"],
            "ticket": z["ticket"], "marcas": z["marcas"], "top1": z["concentracao"]["top1"],
            "top5": z["concentracao"]["top5"], "top10": z["concentracao"]["top10"],
            "crescendo": z["crescendo"], "estavel": z["estavel"], "diminuindo": z["diminuindo"],
            "sat_baixa": z["sat_baixa"], "sat_media": z["sat_media"], "sat_alta": z["sat_alta"],
            "vendedores_mediana": z["vendedores_mediana"],
            "vpv_mediana": statistics.median(vpv) if vpv else 0,
            "novas": z.get("novas"), "sairam": z.get("sairam"), "var_vendas": z.get("var_vendas"),
            "var_unidades": z.get("var_unidades")})
        ant = linhas

    totais = [m["vendas"] for m in resumo_meses]
    por_marca = {}
    for i, linhas in enumerate(linhas_por_mes):
        for l in linhas:
            d = por_marca.setdefault(l["marca_chave"], {"marca": l["marca"], "l": [None] * n})
            d["l"][i] = l
            d["marca"] = l["marca"]          # grafia do mês mais recente

    ult = linhas_por_mes[-1] if linhas_por_mes else []
    vend_ult = sorted(l["vendas"] or 0 for l in ult)
    p75 = vend_ult[int(len(vend_ult) * 0.75)] if vend_ult else 1
    explorador = {nubi.compacta(m): m for m in marcas_explorador}
    k3 = max(0, n - 4)                         # "3 meses atrás" (ou o primeiro)
    marcas = []
    for chave, d in por_marca.items():
        ls = d["l"]
        vendas = [l["vendas"] if l else None for l in ls]
        pos = [l["posicao"] if l else None for l in ls]
        vend = [l["vendedores"] if l else None for l in ls]
        pres = [i for i in range(n) if ls[i]]
        prim, ultm = pres[0], pres[-1]
        atual = ls[-1]
        comp = [(vendas[i - 1], vendas[i]) for i in range(1, n) if ls[i] and ls[i - 1]]
        subidas = sum(1 for a, b in comp if (b or 0) > (a or 0))
        mom = [_cresc(a, b) for a, b in comp if a]
        cresc_total = _cresc(vendas[prim], vendas[-1]) if atual and prim < n - 1 else None
        base3 = next((i for i in range(k3, n - 1) if ls[i]), None)
        cresc_3m = _cresc(vendas[base3], vendas[-1]) if atual and base3 is not None else None
        cresc_mes = _cresc(vendas[-2], vendas[-1]) if n > 1 and atual and ls[-2] else None
        var_vend = _cresc(vend[prim], vend[-1]) if atual and prim < n - 1 and vend[prim] else None
        var_ticket = (_cresc(_div(vendas[prim], ls[prim]["unidades"]), _div(vendas[-1], atual["unidades"]))
                      if atual and prim < n - 1 and ls[prim]["unidades"] and atual["unidades"] else None)
        subida_3m = (pos[base3] - pos[-1]) if atual and base3 is not None else None
        consist = _div(subidas, len(comp)) if comp else None
        volat = statistics.pstdev(mom) if len(mom) >= 2 else None
        share = [_div(v, t) if v is not None else None for v, t in zip(vendas, totais)]
        x = {
            "marca": d["marca"], "marca_chave": chave, "vendas": vendas, "posicoes": pos,
            "share": share, "vendedores_serie": vend, "presente": len(pres), "no_ultimo": bool(atual),
            "primeiro_mes": meses[prim][:7], "ultimo_mes": meses[ultm][:7],
            "posicao": atual["posicao"] if atual else None, "posicao_inicio": pos[prim],
            "ganho_posicoes": (pos[prim] - atual["posicao"]) if atual and prim < n - 1 else None,
            "vendas_atual": vendas[-1], "vendas_inicio": vendas[prim],
            "ganho_vendas": (vendas[-1] - vendas[prim]) if atual and prim < n - 1 else None,
            "ganho_3m": (vendas[-1] - vendas[base3]) if atual and base3 is not None else None,
            "ganho_mes": (vendas[-1] - vendas[-2]) if cresc_mes is not None else None,
            "cresc_total": cresc_total, "cresc_3m": cresc_3m, "cresc_mes": cresc_mes,
            "share_atual": share[-1], "var_share": ((share[-1] or 0) - (share[prim] or 0)) if prim < n - 1 else None,
            "meses_em_alta": subidas, "comparacoes": len(comp), "consistencia": consist, "volatilidade": volat,
            "vendedores": atual["vendedores"] if atual else None, "var_vendedores": var_vend,
            "ticket": _div(vendas[-1], atual["unidades"]) if atual else None, "var_ticket": var_ticket,
            "tendencia": atual["tendencia"] if atual else None, "saturacao": atual["saturacao"] if atual else None,
            "explorador": explorador.get(chave),
        }
        # classificação
        if not atual:
            st = "Saiu do ranking"
        elif n >= 3 and prim >= n - 2:
            st = "Nova"
        elif cresc_3m is not None and cresc_3m >= 0.3 and (consist or 0) >= 0.6:
            st = "Subindo forte"
        elif consist is not None and consist >= 0.7 and (cresc_total or 0) > 0:
            st = "Crescimento consistente"
        elif (cresc_3m is not None and cresc_3m <= -0.2) or (consist is not None and consist <= 0.3 and (cresc_total or 0) < 0):
            st = "Em queda"
        elif volat is not None and volat > 0.35:
            st = "Instável"
        else:
            st = "Estável"
        x["status"] = st
        if atual:
            tam = math.sqrt(min(1.0, _div(vendas[-1] or 0, p75)))
            sat = SAT_LIVRE.get((atual["saturacao"] or "").lower(), 0.5)
            dem = _clip(((cresc_total or 0) - (var_vend or 0)) / 0.5) if cresc_total is not None else 0
            x["nota"] = round(100 * tam * (0.3 * _clip((cresc_3m or 0) / 0.5) + 0.2 * (consist or 0)
                                           + 0.15 * _clip((subida_3m or 0) / 10) + 0.2 * sat + 0.15 * dem))
        else:
            x["nota"] = None
        marcas.append(x)
    marcas.sort(key=lambda x: (x["posicao"] is None, x["posicao"] or 0, -(max(v or 0 for v in x["vendas"]))))

    entradas, saidas = movimentos(meses, por_marca, totais)
    for m in resumo_meses:
        es = [e for e in entradas if e["mes"] == m["mes"]]
        ss = [e for e in saidas if e["mes"] == m["mes"]]
        m.update({"entradas": len(es), "saidas": len(ss), "vendas_entradas": sum(e["vendas"] or 0 for e in es),
                  "vendas_saidas": sum(e["vendas"] or 0 for e in ss)})
    r = {"meses": resumo_meses, "marcas": marcas, "entradas": entradas, "saidas": saidas}
    if n:
        a, b = resumo_meses[0], resumo_meses[-1]
        r["resumo"] = {
            "n_meses": n, "vendas_periodo": sum(totais), "unidades_periodo": sum(m["unidades"] for m in resumo_meses),
            "cresc_vendas": _cresc(a["vendas"], b["vendas"]) if n > 1 else None,
            "cresc_unidades": _cresc(a["unidades"], b["unidades"]) if n > 1 else None,
            "cresc_ticket": _cresc(a["ticket"], b["ticket"]) if n > 1 else None,
            "cresc_mensal": ((_div(b["vendas"], a["vendas"]) ** (1 / (n - 1)) - 1) if n > 1 and a["vendas"] and b["vendas"] else None),
            "var_top5": (b["top5"] - a["top5"]) if n > 1 else None,
            "sempre": sum(1 for x in marcas if x["presente"] == n),
            "passaram": len(marcas),
            "novas_media": _div(sum(m["novas"] or 0 for m in resumo_meses[1:]), n - 1) if n > 1 else None,
            "entradas": len(entradas), "saidas": len(saidas),
            "entradas_ficaram": sum(1 for e in entradas if e["ainda_no_ranking"]),
            "melhor_mes": max(resumo_meses, key=lambda m: m["vendas"])["mes"],
            "pior_mes": min(resumo_meses, key=lambda m: m["vendas"])["mes"],
        }
    return r


def movimentos(meses, por_marca, totais):
    """
    Entradas e saídas do ranking (top do relatório), mês a mês, a partir do 2º mês.
    Entrada: a marca aparece no mês e não estava no mês anterior. Guarda o faturamento e a
    posição com que entrou e o que aconteceu depois (ficou? cresceu desde a entrada?).
    "Cresceu" olha o share (a fatia das vendas do ranking) e a posição, não o R$, para não
    confundir crescimento da marca com o mercado inteiro crescendo.
    Saída: estava no mês anterior e não está neste. Guarda o último faturamento e a posição
    antes de sair, o pico enquanto estava no ranking e se voltou depois.
    """
    n = len(meses)
    entradas, saidas = [], []
    for chave, d in por_marca.items():
        ls = d["l"]
        for i in range(1, n):
            if ls[i] and not ls[i - 1]:
                l = ls[i]
                seg = [j for j in range(i, n) if ls[j]]
                fim = i
                while fim + 1 < n and ls[fim + 1]:
                    fim += 1
                pico = max(seg, key=lambda j: ls[j]["vendas"] or 0)
                atual = ls[-1]
                var_sh = (_cresc(_div(l["vendas"], totais[i]), _div(atual["vendas"], totais[-1]))
                          if atual and i < n - 1 else None)
                if i == n - 1:
                    leitura = "Entrou agora"
                elif not atual:
                    leitura = "Entrou e saiu"
                elif (var_sh or 0) >= 0.2 or atual["posicao"] <= l["posicao"] - 5:
                    leitura = "Ficou e cresceu"
                elif (var_sh or 0) <= -0.2 or atual["posicao"] >= l["posicao"] + 5:
                    leitura = "Ficou, mas caindo"
                else:
                    leitura = "Ficou"
                entradas.append({
                    "mes": meses[i][:7], "mes_nome": nome_mes(meses[i]), "marca": d["marca"], "marca_chave": chave,
                    "posicao": l["posicao"], "vendas": l["vendas"], "unidades": l["unidades"],
                    "vendedores": l["vendedores"], "tendencia": l["tendencia"], "saturacao": l["saturacao"],
                    "reentrada": any(ls[:i - 1]), "meses_seguidos": fim - i + 1,
                    "ainda_no_ranking": bool(atual), "posicao_atual": atual["posicao"] if atual else None,
                    "vendas_atual": atual["vendas"] if atual else None,
                    "cresc_desde_entrada": _cresc(l["vendas"], atual["vendas"]) if atual and i < n - 1 else None,
                    "var_share_desde_entrada": var_sh,
                    "posicoes_desde_entrada": (l["posicao"] - atual["posicao"]) if atual and i < n - 1 else None,
                    "pico_vendas": ls[pico]["vendas"], "pico_posicao": min(ls[j]["posicao"] for j in seg),
                    "leitura": leitura})
            if ls[i - 1] and not ls[i]:
                ini = i - 1
                while ini - 1 >= 0 and ls[ini - 1]:
                    ini -= 1
                run = list(range(ini, i))
                pico = max(run, key=lambda j: ls[j]["vendas"] or 0)
                u = ls[i - 1]
                volta = next((j for j in range(i + 1, n) if ls[j]), None)
                queda = _cresc(ls[pico]["vendas"], u["vendas"])
                if volta is not None:
                    leitura = "Voltou depois"
                elif queda is not None and queda <= -0.3:
                    leitura = "Perdeu fôlego"
                elif len(run) == 1:
                    leitura = "Passou rápido"
                else:
                    leitura = "Saiu"
                saidas.append({
                    "mes": meses[i][:7], "mes_nome": nome_mes(meses[i]), "marca": d["marca"], "marca_chave": chave,
                    "ultimo_mes": meses[i - 1][:7], "posicao": u["posicao"], "vendas": u["vendas"],
                    "unidades": u["unidades"], "vendedores": u["vendedores"], "tendencia": u["tendencia"],
                    "meses_no_ranking": len(run), "desde": meses[ini][:7] if ini > 0 else None,
                    "pico_vendas": ls[pico]["vendas"], "pico_posicao": min(ls[j]["posicao"] for j in run),
                    "pico_mes": meses[pico][:7], "queda_desde_pico": queda,
                    "voltou_em": meses[volta][:7] if volta is not None else None, "leitura": leitura})
    entradas.sort(key=lambda e: (e["mes"], -(e["vendas"] or 0)))
    saidas.sort(key=lambda e: (e["mes"], -(e["vendas"] or 0)))
    return entradas, saidas
